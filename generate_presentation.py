import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme colors
    PRIMARY = RGBColor(15, 23, 42)       # Slate 900 (Deep Navy)
    SECONDARY = RGBColor(37, 99, 235)    # Blue 600 (Vibrant Brand Blue)
    ACCENT = RGBColor(14, 165, 233)      # Sky 500 (Cyan Accent)
    CARD_BG = RGBColor(241, 245, 249)    # Slate 100
    TEXT_DARK = RGBColor(30, 41, 59)     # Slate 800
    TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
    WHITE = RGBColor(255, 255, 255)
    GREEN = RGBColor(16, 185, 129)      # Emerald 500

    def add_header(slide, title_text, category="BATCH-7 INTERNSHIP CLOSURE | COASTALSEVEN TECHNOLOGIES"):
        # Header banner shape
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf = cat_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = SECONDARY

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf2 = title_box.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(26)
        p2.font.bold = True
        p2.font.color.rgb = PRIMARY

    def add_card(slide, left, top, width, height, title, items, badge=""):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = RGBColor(226, 232, 240)
        shape.line.width = Pt(1)

        tx_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True

        if badge:
            p_badge = tf.paragraphs[0]
            p_badge.text = badge.upper()
            p_badge.font.size = Pt(10)
            p_badge.font.bold = True
            p_badge.font.color.rgb = SECONDARY
            p_title = tf.add_paragraph()
        else:
            p_title = tf.paragraphs[0]

        p_title.text = title
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = PRIMARY
        p_title.space_after = Pt(10)

        for item in items:
            p_item = tf.add_paragraph()
            p_item.text = f"•  {item}"
            p_item.font.size = Pt(13)
            p_item.font.color.rgb = TEXT_DARK
            p_item.space_after = Pt(6)

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = PRIMARY
    bg1.line.fill.background()

    # Accent top border
    top_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = SECONDARY
    top_bar.line.fill.background()

    title_box = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(3.0))
    tf1 = title_box.text_frame
    tf1.word_wrap = True
    
    p_tag = tf1.paragraphs[0]
    p_tag.text = "COASTALSEVEN TECHNOLOGIES • KNOWLEDGE FACTORY PROGRAMME"
    p_tag.font.size = Pt(13)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT
    p_tag.space_after = Pt(14)

    p_main = tf1.add_paragraph()
    p_main.text = "My Internship Journey: Transforming Into a Full-Stack & AI Cloud Engineer"
    p_main.font.size = Pt(32)
    p_main.font.bold = True
    p_main.font.color.rgb = WHITE
    p_main.space_after = Pt(14)

    p_sub = tf1.add_paragraph()
    p_sub.text = "A 24-Day Immersion in Full-Stack Web, Generative AI, Microservices & AWS DevOps"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = RGBColor(148, 163, 184)

    # Intern Meta box
    meta_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(5.2), Inches(10.9), Inches(1.4))
    meta_box.fill.solid()
    meta_box.fill.fore_color.rgb = RGBColor(30, 41, 59)
    meta_box.line.color.rgb = RGBColor(51, 65, 85)

    mtf = meta_box.text_frame
    mtf.word_wrap = True
    mp1 = mtf.paragraphs[0]
    mp1.text = "PRESENTED BY:"
    mp1.font.size = Pt(10)
    mp1.font.bold = True
    mp1.font.color.rgb = ACCENT

    mp2 = mtf.add_paragraph()
    mp2.text = "Makkena Thriveni  |  Batch-7 Intern  |  Date: 16th October 2026"
    mp2.font.size = Pt(17)
    mp2.font.bold = True
    mp2.font.color.rgb = WHITE

    s1.notes_slide.notes_text_frame.text = (
        "Good morning mentors, coordinators, and fellow interns! My name is Makkena Thriveni. "
        "Today, I am proud and excited to present my journey in the Batch-7 Knowledge Factory Programme "
        "at CoastalSeven Technologies. Over these 24 intensive days, I evolved from classroom theory to architecting, "
        "building, and deploying production-grade, AI-powered cloud applications."
    )

    # ==========================================
    # SLIDE 2: Expectations vs. Experience
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Expectations vs. Reality: A Mindset Shift")

    add_card(
        s2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "What I Initially Expected",
        [
            "Standard academic project or textbook assignments.",
            "Isolated coding tasks with simple CRUD operations.",
            "Running applications only on localhost:8000.",
            "Basic manual testing without automated quality gates.",
            "Theoretical exposure to AI and cloud concepts without real implementation."
        ],
        badge="Initial Perspective"
    )

    add_card(
        s2, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0),
        "What CoastalSeven Actually Delivered",
        [
            "Enterprise Engineering: Building 'ShopZone' e-commerce platform.",
            "Full Stack Synergy: FastAPI backend + React 19 / Vite frontend.",
            "Real AI in Production: RAG semantic search, ChromaDB vectors, Gemini streaming.",
            "Async Architecture: Celery worker queues, Redis cache & WebSockets.",
            "True DevOps: Docker Compose, Nginx reverse proxy, AWS EC2, and Vercel edge deployment."
        ],
        badge="The Reality & Transformation"
    )

    s2.notes_slide.notes_text_frame.text = (
        "When I joined CoastalSeven, I anticipated typical academic exercises. However, CoastalSeven exceeded all expectations. "
        "Instead of toy apps, we engineered ShopZone—a real, multi-tier platform with AI recommendation engines, Redis caching, "
        "Celery queues, and automated test pipelines. It shifted my perspective from just writing code to building reliable systems."
    )

    # ==========================================
    # SLIDE 3: Internship Journey & Key Activities
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "24-Day Journey: Milestones & Architecture")

    # 4 horizontal progression cards
    add_card(
        s3, Inches(0.8), Inches(1.8), Inches(2.7), Inches(5.0),
        "Phase 1: Backend",
        [
            "FastAPI architecture.",
            "PostgreSQL & SQLite dual engine.",
            "JWT auth & RBAC permissions.",
            "Pydantic schema validation.",
            "Secure password hashing."
        ],
        badge="Days 1 - 8"
    )

    add_card(
        s3, Inches(3.8), Inches(1.8), Inches(2.7), Inches(5.0),
        "Phase 2: Frontend",
        [
            "React 19 + Vite 8 SPA.",
            "Zustand state management.",
            "Responsive shopping bag & orders.",
            "Live WebSocket chat widget.",
            "Dark & Light mode themes."
        ],
        badge="Days 9 - 15"
    )

    add_card(
        s3, Inches(6.8), Inches(1.8), Inches(2.7), Inches(5.0),
        "Phase 3: AI & Scale",
        [
            "Gemini Generative AI.",
            "RAG Semantic Search.",
            "ChromaDB vector database.",
            "Celery async background tasks.",
            "Redis 7 caching & broker."
        ],
        badge="Days 16 - 22"
    )

    add_card(
        s3, Inches(9.8), Inches(1.8), Inches(2.7), Inches(5.0),
        "Phase 4: Cloud & DevOps",
        [
            "Docker multi-container stack.",
            "AWS EC2 Linux deployment.",
            "Nginx SSL / Let's Encrypt.",
            "Sentry & Prometheus metrics.",
            "Vercel multi-service hosting."
        ],
        badge="Days 23 - 24"
    )

    s3.notes_slide.notes_text_frame.text = (
        "Our journey was structured in four progressive phases: Starting with robust FastAPI backend foundations and security, "
        "advancing to modern React 19 UI with real-time WebSockets, integrating GenAI and RAG with ChromaDB and Celery for async tasks, "
        "and concluding with complete DevOps—Docker containerization, AWS EC2, Nginx HTTPS, and Vercel deployment."
    )

    # ==========================================
    # SLIDE 4: Technical Skills & Knowledge Gained
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Technical Skills & Competency Matrix")

    add_card(
        s4, Inches(0.8), Inches(1.8), Inches(3.7), Inches(5.0),
        "Full-Stack Development",
        [
            "FastAPI & Python 3.14: Async endpoints, middleware, dependency injection.",
            "React 19 & TypeScript: Component lifecycle, hooks, and responsive UX.",
            "PostgreSQL & SQLite: Complex relational queries, foreign keys, migrations.",
            "Real-Time WebSockets: Bi-directional communication for live support chat."
        ],
        badge="Application Core"
    )

    add_card(
        s4, Inches(4.8), Inches(1.8), Inches(3.7), Inches(5.0),
        "AI & Asynchronous Processing",
        [
            "RAG Architecture: Retrieving relevant catalog context for LLMs.",
            "ChromaDB: Embedding storage, similarity search, and index management.",
            "Celery & Redis: Background invoice PDF generation & catalog reindexing.",
            "Redis AI Cache: Sub-millisecond response caching for chat queries."
        ],
        badge="Artificial Intelligence"
    )

    add_card(
        s4, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
        "Cloud, DevOps & Quality",
        [
            "Docker & Compose: Multi-container orchestration with health checks.",
            "AWS EC2 & IAM: Security groups, least privilege, and SSH hardening.",
            "Observability: Sentry error tracking & Prometheus /metrics.",
            "Test-Driven Engineering: 132 automated tests with 100% pass rate."
        ],
        badge="Production & Cloud"
    )

    s4.notes_slide.notes_text_frame.text = (
        "Here is the technical matrix I've mastered. On the core stack: FastAPI, React 19, and TypeScript. "
        "On AI: Retrieval-Augmented Generation with ChromaDB, paired with Celery distributed workers and Redis caching. "
        "On cloud and DevOps: Docker multi-containerization, AWS EC2, Nginx SSL, Prometheus metrics, and test automation."
    )

    # ==========================================
    # SLIDE 5: Personal Growth & Achievements
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Key Achievements & Professional Growth")

    # Metrics Row
    m1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(3.6), Inches(1.3))
    m1.fill.solid()
    m1.fill.fore_color.rgb = CARD_BG
    m1.line.color.rgb = SECONDARY
    m1.line.width = Pt(2)
    tf_m1 = m1.text_frame
    tf_m1.word_wrap = True
    p = tf_m1.paragraphs[0]
    p.text = "132 / 132 TESTS"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = SECONDARY
    p2 = tf_m1.add_paragraph()
    p2.text = "100% Pass Rate across Unit, Integration & UI suites"
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_DARK

    m2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(1.8), Inches(3.7), Inches(1.3))
    m2.fill.solid()
    m2.fill.fore_color.rgb = CARD_BG
    m2.line.color.rgb = GREEN
    m2.line.width = Pt(2)
    tf_m2 = m2.text_frame
    tf_m2.word_wrap = True
    p = tf_m2.paragraphs[0]
    p.text = "COMPLETE SHOPZONE"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = GREEN
    p2 = tf_m2.add_paragraph()
    p2.text = "End-to-End E-Commerce Platform Built from Scratch"
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_DARK

    m3 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(1.8), Inches(3.7), Inches(1.3))
    m3.fill.solid()
    m3.fill.fore_color.rgb = CARD_BG
    m3.line.color.rgb = ACCENT
    m3.line.width = Pt(2)
    tf_m3 = m3.text_frame
    tf_m3.word_wrap = True
    p = tf_m3.paragraphs[0]
    p.text = "CLOUD DEPLOYED"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = ACCENT
    p2 = tf_m3.add_paragraph()
    p2.text = "Live Docker Compose, AWS EC2 & Vercel Services"
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_DARK

    # Growth bullets
    add_card(
        s5, Inches(0.8), Inches(3.4), Inches(11.7), Inches(3.5),
        "Core Growth Highlights",
        [
            "From Student to Software Engineer: Overcame the fear of debugging complex system errors, asynchronous race conditions, and container networking.",
            "Architectural Thinking: Learned to design for scalability—separating heavy tasks into Celery queues and caching expensive queries in Redis.",
            "Security-First Mindset: Implemented PII redaction, JWT token expiry, CORS policies, and IAM least privilege rather than leaving credentials exposed.",
            "Clean Code & Resilience: Built structured JSON logging with correlation IDs and deep health checks to ensure production reliability."
        ],
        badge="Mindset & Capabilities"
    )

    s5.notes_slide.notes_text_frame.text = (
        "Beyond the code, the personal transformation has been immense. We built a fully working, tested platform backed by 132 passing tests. "
        "More importantly, I developed architectural thinking: understanding why we need message brokers, how to secure microservices, "
        "and how to debug distributed containers. I feel confident approaching complex production software challenges."
    )

    # ==========================================
    # SLIDE 6: Feedback, Takeaways & Gratitude
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Takeaways, Suggestions & Gratitude")

    add_card(
        s6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Key Takeaways & Experience Feedback",
        [
            "Practical, Hands-on Learning: The Knowledge Factory model of building real features daily made theoretical concepts stick immediately.",
            "Industry Standards: Working with Docker, Sentry, Redis, and Git CI/CD mirrors how world-class tech companies operate.",
            "Constructive Mentorship: The guidance from CoastalSeven mentors helped troubleshoot bottlenecks and elevated code quality.",
            "Suggestion: Adding an open-ended collaborative hackathon day would be an exciting addition for future batches!"
        ],
        badge="Reflections"
    )

    add_card(
        s6, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0),
        "Heartfelt Gratitude",
        [
            "To CoastalSeven Technologies: Thank you for organizing the Knowledge Factory Programme and investing in our growth.",
            "To Our Mentors & Leaders: Thank you for your patience, daily code reviews, technical insights, and encouragement.",
            "To Fellow Batch-7 Interns: Thank you for the inspiring discussions, peer support, and camaraderie throughout this journey.",
            "Next Steps: Ready to leverage this full-stack, AI, and cloud expertise in my professional career ahead!"
        ],
        badge="Appreciation"
    )

    s6.notes_slide.notes_text_frame.text = (
        "I want to express my deepest gratitude to Team CoastalSeven Technologies and our dedicated mentors. "
        "The Knowledge Factory Programme was rigorous, practical, and truly transformative. It gave me the skills, "
        "confidence, and mindset needed to thrive in modern software engineering. Thank you for this wonderful opportunity!"
    )

    # ==========================================
    # SLIDE 7: Conclusion & Q&A
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    bg7 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg7.fill.solid()
    bg7.fill.fore_color.rgb = PRIMARY
    bg7.line.fill.background()

    top_bar7 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.15))
    top_bar7.fill.solid()
    top_bar7.fill.fore_color.rgb = GREEN
    top_bar7.line.fill.background()

    q_box = s7.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.3), Inches(3.5))
    tf7 = q_box.text_frame
    tf7.word_wrap = True

    qp1 = tf7.paragraphs[0]
    qp1.text = "THANK YOU!"
    qp1.font.size = Pt(44)
    qp1.font.bold = True
    qp1.font.color.rgb = WHITE
    qp1.alignment = PP_ALIGN.CENTER
    qp1.space_after = Pt(16)

    qp2 = tf7.add_paragraph()
    qp2.text = "CoastalSeven Technologies • Knowledge Factory Programme • Batch-7"
    qp2.font.size = Pt(18)
    qp2.font.color.rgb = ACCENT
    qp2.alignment = PP_ALIGN.CENTER
    qp2.space_after = Pt(24)

    qp3 = tf7.add_paragraph()
    qp3.text = "Makkena Thriveni  |  Full-Stack & Cloud Engineering Intern"
    qp3.font.size = Pt(16)
    qp3.font.color.rgb = RGBColor(226, 232, 240)
    qp3.alignment = PP_ALIGN.CENTER
    qp3.space_after = Pt(16)

    qp4 = tf7.add_paragraph()
    qp4.text = "Questions & Discussion"
    qp4.font.size = Pt(22)
    qp4.font.bold = True
    qp4.font.color.rgb = GREEN
    qp4.alignment = PP_ALIGN.CENTER

    s7.notes_slide.notes_text_frame.text = (
        "Thank you all for your time and attention. I am now open to any questions, thoughts, or feedback!"
    )

    output_path = r"C:\Users\Dell\Downloads\day9-fastapi\Makkena_Thriveni_Internship_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation()
